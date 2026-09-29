# /// script
# requires-python = ">=3.11"
# dependencies = ["python-docx>=1.1", "jsonschema>=4.0"]
# ///
"""
La BITÁCORA del grupo (épico #26, ADR-0005): una plantilla .docx lista para llenar (sin pre-llenar)
y su conversión a dato estructurado conforme a data/schema/bitacora.schema.json.

Una sola definición de la estructura (SECCIONES y compañía, abajo) sirve para GENERAR la plantilla,
LLENARLA (ejemplo / pruebas) y LEERLA. El lector reconoce cada tabla por el texto de su primera
celda y cada fila por su etiqueta, no por su posición: tolera filas agregadas o vacías, párrafos
múltiples, mayúsculas y tildes (ediciones típicas en Word o Google Docs).

Convención de celdas (igual que el schema):
- texto            -> lo que escribió el grupo
- «Sin dato»       -> null: HUECO de información declarado (la fuente no lo dice; es un hallazgo)
- vacía            -> no llenada: se omite y se lista en `sin_llenar`

Uso (desde la raíz):
  uv run extraccion/bitacora.py generar                        # -> web/public/bitacora-laboratorio.docx
  uv run extraccion/bitacora.py ejemplo                        # -> web/public/bitacora-ejemplo-ctei.docx
  uv run extraccion/bitacora.py leer grupo3.docx --slug ciencia-tecnologia
                                                               # -> data/bitacoras/<slug>/<grupo>.json
  uv run extraccion/bitacora.py probar                         # ida y vuelta con el ejemplo CTeI
"""
from __future__ import annotations

import argparse
import json
import re
import sys
import tempfile
import unicodedata
from dataclasses import dataclass
from pathlib import Path

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor
from jsonschema import Draft202012Validator

VERSION = 4  # la v4 de la plantilla que iteró el equipo; se imprime en el pie
VIGENCIAS = ["2018-2022", "2022-2026"]
ETIQUETA_VIGENCIA = {"2018-2022": "2018–2022 (Duque)", "2022-2026": "2022–2026 (Petro)"}
SIN_DATO = "Sin dato"

SCHEMA = Path("data/schema/bitacora.schema.json")
PLANTILLA = Path("web/public/bitacora-laboratorio.docx")
EJEMPLO_DOCX = Path("web/public/bitacora-ejemplo-ctei.docx")
EJEMPLO_JSON = Path(__file__).parent / "ejemplos" / "bitacora-ctei.json"
DATOS_WEB = Path("web/src/lib/data")
SALIDA = Path("data/bitacoras")


# ─────────────────────────── La estructura (definición única) ───────────────────────────


@dataclass(frozen=True)
class Fila:
    clave: str
    etiqueta: str
    pista: str = ""


@dataclass(frozen=True)
class Comparativa:
    """Tabla lado a lado: una fila por campo, una columna por gobierno."""

    clave: str  # clave en el JSON: {clave: {vigencia: {fila: celda}}}
    titulo: str  # encabezado de la sección
    cabecera: str  # texto de la primera celda: identifica la tabla al leer
    guia: str
    filas: tuple[Fila, ...]


@dataclass(frozen=True)
class Caja:
    """Texto libre: tabla de una columna (título y pregunta + espacio para escribir)."""

    clave: str
    titulo: str  # primera línea de la cabecera: identifica la tabla al leer
    pista: str
    alto_cm: float = 4.0


@dataclass(frozen=True)
class Lista:
    """Tabla de registros: una fila por registro, columnas fijas. Se leen solo las filas con algo
    escrito; el grupo puede agregar filas."""

    clave: str
    cabecera: tuple[str, ...]  # la primera celda identifica la tabla al leer
    claves: tuple[str, ...]  # clave en el JSON de cada columna
    anchos_cm: tuple[float, ...]


DATOS = (
    Fila("grupo", "Grupo", "Nombre o número"),
    Fila("integrantes", "Integrantes", "Nombres, separados por coma"),
    Fila("politica", "Política pública", "Como aparece en la red"),
    Fila("motivo", "¿Por qué elegimos esta política?", "Una pregunta o una razón concreta"),
    Fila("expectativa", "¿Qué esperamos encontrar?", "Nuestra hipótesis inicial, antes de comparar"),
)

# 1. Ubicar la política
UBICACION = Comparativa(
    "ubicacion",
    "1. Ubicar la política",
    "Qué mirar",
    "Ubiquemos la política en el informe de cada gobierno y veamos cómo reporta su avance. La red puede "
    "sugerir dónde está, pero verifiquémoslo en el informe.",
    (
        Fila("ubicacion", "Ubicación", "Sección, capítulo o páginas"),
        Fila("indicador", "Indicador o forma de reporte", "¿Cómo muestra el avance? ¿Qué mide?"),
        Fila("valor", "Valor y referencia", "Cifra reportada y su meta, periodo o base de comparación"),
    ),
)
COMPARABILIDAD = Caja(
    "comparabilidad",
    "Tabla puente: ¿qué es comparable?",
    "¿Los dos gobiernos miden lo mismo, con la misma unidad, la misma población y la misma meta? ¿Qué no "
    "se puede comparar?",
    4.5,
)

# 2. Los instrumentos: una fila por instrumento (hasta tres), una columna por gobierno.
INSTRUMENTOS_TITULO = "2. Los instrumentos"
INSTRUMENTOS_GUIA = (
    "Escojamos hasta tres instrumentos de esta política: el medio concreto con que actúa el Estado (un "
    "programa, una ley, un fondo, una convocatoria). Describamos cada uno en los dos gobiernos: ¿sigue, "
    "cambia de uso, es nuevo o se deja? En la última columna, digamos por qué lo elegimos."
)
INSTRUMENTOS_CABECERA = "Instrumento"
INSTRUMENTOS_RELEVANCIA = "Por qué es relevante"
INSTRUMENTOS_FILAS = 3
INSTRUMENTOS_NOMBRE = "Nombre:"  # pista de la celda de la etiqueta: el nombre va a continuación

# 3. Co-construyamos la red
APORTES_TITULO = "3. Co-construyamos la red"
APORTES_GUIA = (
    "La red es una primera propuesta. Corrijamos, añadamos, conectemos o reclasifiquemos lo que haga falta, "
    "y anotemos las dudas. Si hay varios aportes del mismo tipo, agreguemos filas."
)
APORTES = Lista(
    "aportes",
    ("Tipo de aporte", "Elemento o relación", "Nuestra propuesta", "Evidencia o fuente"),
    ("tipo", "elemento", "propuesta", "evidencia"),
    (3.6, 4.4, 5.4, 4.2),
)
TIPOS_APORTE = (
    Fila("corregir", "Corregir", "Un dato de la red que está mal"),
    Fila("anadir", "Añadir", "Un instrumento o una política que falta"),
    Fila("conectar", "Conectar", "Una relación que falta"),
    Fila("reclasificar", "Reclasificar", "Le corresponde otro tipo de instrumento o de cambio"),
    Fila("duda", "Duda", "Algo que hay que revisar"),
    Fila("otro", "Otro"),
)

# 4. Las siete subcategorías
SUBCATEGORIAS = Comparativa(
    "subcategorias",
    "4. Las siete subcategorías, lado a lado",
    "Subcategoría",
    "Comparemos las siete subcategorías, fila por fila. Si la fuente no lo dice, escribamos «Sin dato».",
    (
        Fila("objetivo", "4.1 Objetivo", "¿Qué declara que busca?"),
        Fila("instituciones", "4.2 Instituciones", "¿Quién la ejecuta o la coordina?"),
        Fila("poblacion", "4.3 Población", "¿A quién va dirigida? ¿Atiende distinto a ciertos grupos o regiones?"),
        Fila("normativa", "4.4 Normativa", "¿Qué leyes, decretos o documentos CONPES la sustentan?"),
        Fila("recursos", "4.5 Recursos", "¿Qué recursos reporta y de qué fuente?"),
        Fila("metas", "4.6 Metas", "¿Qué meta se fijó para el cuatrienio y cuánto cumplió?"),
        Fila("impacto", "4.7 Impacto", "¿Hay pruebas de que cambió la vida de la gente, o solo de lo que se hizo y se entregó?"),
    ),
)
COMPARATIVAS = (UBICACION, SUBCATEGORIAS)

# 5. Fuentes complementarias
FUENTES_TITULO = "5. Fuentes complementarias"
FUENTES_GUIA = (
    "Anotemos solo las fuentes distintas del informe que usamos para verificar o ampliar la comparación. Si "
    "hace falta, agreguemos filas."
)
FUENTES = Lista("fuentes", ("Fuente", "Qué aportó", "Página o enlace"), ("fuente", "aporte", "referencia"), (6.0, 6.6, 5.0))
FUENTES_FILAS = 4

# 6. En términos de complejidad
COMPLEJIDAD_TITULO = "6. En términos de complejidad"
COMPLEJIDAD = Caja(
    "complejidad",
    "Patrones que pueden emerger",
    "¿Qué podría verse al mirar juntas varias políticas que no se ve en una sola? Por ejemplo, instrumentos, "
    "entidades o recursos que comparten.",
)

# 7. Conclusiones
CONCLUSIONES_TITULO = "7. Conclusiones"
HALLAZGOS = Caja("hallazgos", "Hallazgo principal", "¿Qué aprendimos que no era evidente al comienzo?", 5.0)
HIPOTESIS = Caja("hipotesis", "Hipótesis para la sesión 2", "¿Qué patrón creemos que podría repetirse en otras políticas?")
CAJAS = (COMPARABILIDAD, COMPLEJIDAD, HALLAZGOS, HIPOTESIS)


# ─────────────────────────── Utilidades ───────────────────────────


def normalizar(s: str, *, numeracion: bool = True) -> str:
    """Sin tildes, minúsculas, espacios colapsados; por defecto sin numeración inicial."""
    s = unicodedata.normalize("NFD", s)
    s = "".join(c for c in s if not unicodedata.combining(c)).lower()
    s = s.replace("–", "-").replace("—", "-")
    if numeracion:
        s = re.sub(r"^\s*\d+(\s*[.\-]\s*\d+)*\.?\s+", "", s)  # «2. », «1.1 », «2-8. »
    return re.sub(r"\s+", " ", s).strip(" :")


def slug(s: str) -> str:
    return re.sub(r"[^a-z0-9]+", "-", normalizar(s)).strip("-") or "sin-nombre"


def _texto_celda(celda) -> str:
    """Todo el texto de la celda, incluidas tablas pegadas dentro (se leen fila por fila)."""
    lineas = [p.text.strip() for p in celda.paragraphs]
    for anidada in celda.tables:
        for fila in anidada.rows:
            lineas.append(" · ".join(t for t in (_texto_celda(c) for c in fila.cells) if t))
    return "\n".join(l for l in lineas if l)


def _primera_linea(celda) -> str:
    for p in celda.paragraphs:
        if p.text.strip():
            return p.text.strip()
    return ""


_HUECO = re.compile(r"^\s*(sin\s+datos?|s\s*/\s*d|no\s+hay\s+datos?)\b\.?(.*)$", re.IGNORECASE | re.DOTALL)


def _valor(texto: str):
    """texto -> (presente, valor, nota). «Sin dato» (o «S/D», «No hay dato») es null: hueco
    declarado; «Sin dato: <nota>» o «Sin dato (<nota>)» conserva la nota aparte. Vacío = no llenado."""
    if not texto:
        return False, None, None
    m = _HUECO.match(texto)
    if m:
        nota = m.group(2).strip().lstrip(":-–—(").strip().rstrip(".)").strip()
        return True, None, nota or None
    return True, texto, None


# ─────────────────────────── Generar ───────────────────────────

GRIS = RGBColor(0x5D, 0x63, 0x71)
TINTA = RGBColor(0x23, 0x26, 0x2E)
MORADO = RGBColor(0x4B, 0x2E, 0x83)  # títulos y cabeceras (paleta de la v4 del equipo)
MORADO_OSCURO = RGBColor(0x3E, 0x1A, 0x63)
LILA = RGBColor(0x69, 0x41, 0xA5)
BLANCO = RGBColor(0xFF, 0xFF, 0xFF)
FONDO_CABECERA = "4B2E83"
FONDO_ETIQUETA = "EEE8F7"
FONDO_ALTERNO = "F7F4FB"
BORDE = "D7CBE7"
ANCHO_UTIL_CM = 17.6  # carta con márgenes de 2 cm


def _sombrear(celda, color: str) -> None:
    tc = celda._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), color)
    tc.append(shd)


def _bordes(tabla) -> None:
    """Bordes finos lila en toda la tabla (en vez del negro de «Table Grid»)."""
    tblPr = tabla._tbl.tblPr
    bordes = OxmlElement("w:tblBorders")
    for lado in ("top", "left", "bottom", "right", "insideH", "insideV"):
        b = OxmlElement(f"w:{lado}")
        b.set(qn("w:val"), "single")
        b.set(qn("w:sz"), "6")
        b.set(qn("w:space"), "0")
        b.set(qn("w:color"), BORDE)
        bordes.append(b)
    tblPr.append(bordes)


def _alto_minimo(fila, cm: float) -> None:
    trPr = fila._tr.get_or_add_trPr()
    alto = OxmlElement("w:trHeight")
    alto.set(qn("w:val"), str(int(cm * 567)))
    alto.set(qn("w:hRule"), "atLeast")
    trPr.append(alto)


def _no_partir(fila) -> None:
    trPr = fila._tr.get_or_add_trPr()
    trPr.insert(0, OxmlElement("w:cantSplit"))  # va antes de trHeight en el esquema


def _cabecera_pegada(fila) -> None:
    """La fila de encabezado se repite en cada página y no queda sola al pie."""
    trPr = fila._tr.get_or_add_trPr()
    trPr.append(OxmlElement("w:tblHeader"))
    for celda in fila.cells:
        for p in celda.paragraphs:
            p.paragraph_format.keep_with_next = True


def _anchos(tabla, anchos_cm) -> None:
    tabla.autofit = False
    tabla.alignment = WD_TABLE_ALIGNMENT.CENTER
    for fila in tabla.rows:
        for celda, ancho in zip(fila.cells, anchos_cm):
            celda.width = Cm(ancho)


def _escribir(celda, texto: str, *, negrita=False, color=None, tam=None, cursiva=False) -> None:
    p = celda.paragraphs[0]
    r = p.add_run(texto)
    r.bold = negrita
    r.italic = cursiva
    if color is not None:
        r.font.color.rgb = color
    if tam is not None:
        r.font.size = Pt(tam)


def _pista(celda, texto: str, color=GRIS) -> None:
    p = celda.add_paragraph()
    r = p.add_run(texto)
    r.italic = True
    r.font.size = Pt(8.5)
    r.font.color.rgb = color


def _etiqueta(celda, fila: Fila) -> None:
    _escribir(celda, fila.etiqueta, negrita=True, color=MORADO_OSCURO)
    if fila.pista:
        _pista(celda, fila.pista)


def _cabecera(fila, textos) -> None:
    """Fila de encabezado: fondo morado, texto blanco en negrita."""
    for celda, texto in zip(fila.cells, textos):
        _escribir(celda, texto, negrita=True, color=BLANCO, tam=9)
        _sombrear(celda, FONDO_CABECERA)
    _cabecera_pegada(fila)


def _guia(doc, texto: str) -> None:
    p = doc.add_paragraph()
    r = p.add_run(texto)
    r.italic = True
    r.font.color.rgb = GRIS
    p.paragraph_format.space_after = Pt(6)
    p.paragraph_format.keep_with_next = True


def _titulo(doc, texto: str) -> None:
    h = doc.add_heading(texto, level=1)
    h.paragraph_format.space_before = Pt(18)
    h.paragraph_format.space_after = Pt(4)
    h.paragraph_format.keep_with_next = True
    for r in h.runs:
        r.font.color.rgb = MORADO
        r.font.size = Pt(14)


def _tabla(doc, filas: int, cols: int):
    t = doc.add_table(rows=filas, cols=cols)
    t.style = "Table Grid"
    _bordes(t)
    return t


def _tabla_comparativa(doc, c: Comparativa) -> None:
    t = _tabla(doc, 1 + len(c.filas), 3)
    _cabecera(t.rows[0], [c.cabecera, *(ETIQUETA_VIGENCIA[v] for v in VIGENCIAS)])
    for fila, def_fila in zip(t.rows[1:], c.filas):
        _etiqueta(fila.cells[0], def_fila)
        _sombrear(fila.cells[0], FONDO_ETIQUETA)
        _alto_minimo(fila, 2.0)
        _no_partir(fila)
    _anchos(t, [4.4, 6.6, 6.6])


def _tabla_instrumentos(doc) -> None:
    t = _tabla(doc, 1 + INSTRUMENTOS_FILAS, 4)
    _cabecera(t.rows[0], [INSTRUMENTOS_CABECERA, *(ETIQUETA_VIGENCIA[v] for v in VIGENCIAS), INSTRUMENTOS_RELEVANCIA])
    for i, fila in enumerate(t.rows[1:], start=1):
        _escribir(fila.cells[0], f"Instrumento {i}", negrita=True, color=MORADO_OSCURO)
        _pista(fila.cells[0], INSTRUMENTOS_NOMBRE)
        _sombrear(fila.cells[0], FONDO_ETIQUETA)
        _alto_minimo(fila, 4.0)
        _no_partir(fila)
    _anchos(t, [3.8, 4.8, 4.8, 4.2])


def _tabla_lista(doc, l: Lista, etiquetas: tuple[Fila, ...] = (), vacias: int = 0) -> None:
    t = _tabla(doc, 1 + (len(etiquetas) or vacias), len(l.cabecera))
    _cabecera(t.rows[0], l.cabecera)
    for i, fila in enumerate(t.rows[1:]):
        if etiquetas:
            _etiqueta(fila.cells[0], etiquetas[i])
            _sombrear(fila.cells[0], FONDO_ETIQUETA)
        elif i % 2:
            for celda in fila.cells:
                _sombrear(celda, FONDO_ALTERNO)
        _alto_minimo(fila, 1.2)
    _anchos(t, l.anchos_cm)


def _caja(doc, c: Caja) -> None:
    t = _tabla(doc, 2, 1)
    cab = t.rows[0].cells[0]
    _escribir(cab, c.titulo, negrita=True, color=BLANCO)
    _pista(cab, c.pista, BLANCO)
    _sombrear(cab, FONDO_CABECERA)
    _cabecera_pegada(t.rows[0])
    _alto_minimo(t.rows[1], c.alto_cm)
    _no_partir(t.rows[1])
    _anchos(t, [ANCHO_UTIL_CM])
    doc.add_paragraph().paragraph_format.space_after = Pt(2)


def _pie(seccion) -> None:
    p = seccion.footer.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run(f"Laboratorio de Cocreación · Bitácora de grupo · v{VERSION} · página ")
    r.font.size = Pt(8)
    r.font.color.rgb = GRIS
    # Campo PAGE (número de página).
    for tipo, texto in (("begin", None), (None, "PAGE"), ("end", None)):
        run = p.add_run()
        run.font.size = Pt(8)
        run.font.color.rgb = GRIS
        if tipo:
            fld = OxmlElement("w:fldChar")
            fld.set(qn("w:fldCharType"), tipo)
            run._r.append(fld)
        else:
            instr = OxmlElement("w:instrText")
            instr.set(qn("xml:space"), "preserve")
            instr.text = texto
            run._r.append(instr)


def plantilla() -> Document:
    doc = Document()
    sec = doc.sections[0]
    sec.page_width, sec.page_height = Cm(21.59), Cm(27.94)  # carta
    for lado in ("left_margin", "right_margin", "top_margin", "bottom_margin"):
        setattr(sec, lado, Cm(2))
    normal = doc.styles["Normal"]
    normal.font.name = "Arial"
    normal.font.size = Pt(10)
    normal.font.color.rgb = TINTA
    _pie(sec)

    t = doc.add_paragraph()
    r = t.add_run("Bitácora de grupo")
    r.bold = True
    r.font.size = Pt(22)
    r.font.color.rgb = MORADO_OSCURO
    s = doc.add_paragraph()
    r = s.add_run("Laboratorio de Cocreación · Una política pública, dos gobiernos")
    r.font.color.rgb = LILA
    r.font.size = Pt(11.5)

    doc.add_paragraph(
        "Elijamos una política pública de la red y describámosla en los dos gobiernos (2018–2022 y 2022–2026). "
        "Usemos el informe de empalme (el balance que un gobierno le entrega al siguiente) y otras fuentes: "
        "Sinergia (el sistema oficial de seguimiento de metas), el Plan Nacional de Desarrollo, las leyes y los "
        "documentos CONPES (las políticas que aprueba el consejo de planeación)."
    )
    for norma, detalle in (
        (
            "Regla 1: si la fuente no lo dice, escribamos «Sin dato». ",
            "Podemos agregar el porqué: «Sin dato: el balance no lo desagrega». Un vacío es un hallazgo: no lo "
            "llenemos con suposiciones.",
        ),
        (
            "Regla 2: no cambiemos los títulos de tablas y filas ni combinemos celdas. ",
            "Así el equipo organizador puede sumar todas las bitácoras a la red.",
        ),
    ):
        regla = doc.add_paragraph()
        regla.add_run(norma).bold = True
        regla.add_run(detalle)

    datos = _tabla(doc, len(DATOS), 2)
    for fila, d in zip(datos.rows, DATOS):
        _etiqueta(fila.cells[0], d)
        _sombrear(fila.cells[0], FONDO_ETIQUETA)
        _alto_minimo(fila, 1.0)
    _anchos(datos, [6.0, ANCHO_UTIL_CM - 6.0])

    _titulo(doc, UBICACION.titulo)
    _guia(doc, UBICACION.guia)
    _tabla_comparativa(doc, UBICACION)
    doc.add_paragraph().paragraph_format.space_after = Pt(2)
    _caja(doc, COMPARABILIDAD)

    _titulo(doc, INSTRUMENTOS_TITULO)
    _guia(doc, INSTRUMENTOS_GUIA)
    _tabla_instrumentos(doc)

    _titulo(doc, APORTES_TITULO)
    _guia(doc, APORTES_GUIA)
    _tabla_lista(doc, APORTES, TIPOS_APORTE)

    _titulo(doc, SUBCATEGORIAS.titulo)
    _guia(doc, SUBCATEGORIAS.guia)
    _tabla_comparativa(doc, SUBCATEGORIAS)

    _titulo(doc, FUENTES_TITULO)
    _guia(doc, FUENTES_GUIA)
    _tabla_lista(doc, FUENTES, vacias=FUENTES_FILAS)

    _titulo(doc, COMPLEJIDAD_TITULO)
    _caja(doc, COMPLEJIDAD)

    _titulo(doc, CONCLUSIONES_TITULO)
    _caja(doc, HALLAZGOS)
    _caja(doc, HIPOTESIS)
    return doc


# ─────────────────────────── Llenar (ejemplo y pruebas) ───────────────────────────


def _poner(celda, valor) -> None:
    """string -> texto (una línea por párrafo); None -> «Sin dato»."""
    lineas = [SIN_DATO] if valor is None else str(valor).split("\n")
    celda.paragraphs[0].text = lineas[0]
    for l in lineas[1:]:
        celda.add_paragraph(l)


def _agregar(celda, valor) -> None:
    """Como _poner, pero debajo de lo que ya tiene la celda (la etiqueta y su pista)."""
    for l in [SIN_DATO] if valor is None else str(valor).split("\n"):
        celda.add_paragraph(l)


def _filas_para(t, n: int) -> list:
    """Filas de datos de la tabla, agregando las que falten para n registros."""
    while len(t.rows) - 1 < n:
        t.add_row()
    return list(t.rows[1:])


def llenar(doc: Document, b: dict) -> Document:
    tablas = _indexar(doc)
    valores = {
        "grupo": b["grupo"]["nombre"],
        "integrantes": ", ".join(b["grupo"]["integrantes"]),
        "politica": b["politica"]["nombre"],
        "motivo": b["grupo"].get("motivo", ""),
        "expectativa": b["grupo"].get("expectativa", ""),
    }
    for fila in tablas["datos"].rows:
        clave = _fila_de(DATOS, _primera_linea(fila.cells[0]))
        if clave and valores.get(clave):
            _poner(fila.cells[1], valores[clave])
    for c in COMPARATIVAS:
        t = tablas[c.clave]
        cols = _columnas_vigencia(t)
        for fila in t.rows[1:]:
            clave = _fila_de(c.filas, _primera_linea(fila.cells[0]))
            for v, i in cols.items():
                if clave and clave in b[c.clave][v]:
                    _poner(fila.cells[i], b[c.clave][v][clave])
    for fila, ins in zip(tablas["instrumentos"].rows[1:], b["instrumentos"]):
        if "nombre" in ins:
            _agregar(fila.cells[0], ins["nombre"])
        for i, k in enumerate([*VIGENCIAS, "relevancia"], start=1):
            if k in ins:
                _poner(fila.cells[i], ins[k])
    for aporte in b["aportes"]:  # cada aporte en la fila de su tipo; si ya está ocupada, una fila nueva
        t = tablas["aportes"]
        libre = next(
            (f for f in t.rows[1:] if _fila_de(TIPOS_APORTE, _primera_linea(f.cells[0])) == aporte["tipo"]
             and not any(_texto_celda(x) for x in f.cells[1:])),
            None,
        )
        if libre is None:
            libre = t.add_row()
            libre.cells[0].paragraphs[0].text = next(f.etiqueta for f in TIPOS_APORTE if f.clave == aporte["tipo"])
        for celda, k in zip(libre.cells[1:], APORTES.claves[1:]):
            if aporte.get(k):
                _poner(celda, aporte[k])
    for c in CAJAS:
        if c.clave in b:
            _poner(tablas[c.clave].rows[1].cells[0], b[c.clave])
    for fila, fuente in zip(_filas_para(tablas["fuentes"], len(b["fuentes"])), b["fuentes"]):
        for celda, k in zip(fila.cells, FUENTES.claves):
            if fuente.get(k):
                _poner(celda, fuente[k])
    return doc


# ─────────────────────────── Leer ───────────────────────────


class PlantillaInvalida(Exception):
    pass


def _fila_de(filas: tuple[Fila, ...], etiqueta: str) -> str | None:
    """Etiqueta exacta, o ampliada por el grupo en la misma línea («Objetivo del gobierno»)."""
    n = normalizar(etiqueta)
    for f in filas:
        if n == normalizar(f.etiqueta):
            return f.clave
    ampliadas = [f for f in filas if n.startswith(normalizar(f.etiqueta) + " ")]
    return max(ampliadas, key=lambda f: len(f.etiqueta)).clave if ampliadas else None


def _indexar(doc: Document) -> dict:
    """Clave de sección -> tabla, reconocida por el texto de su primera celda."""
    por_cabecera = {normalizar(c.cabecera): c.clave for c in COMPARATIVAS}
    por_cabecera |= {normalizar(c.titulo): c.clave for c in CAJAS}
    por_cabecera[normalizar(INSTRUMENTOS_CABECERA)] = "instrumentos"
    por_cabecera[normalizar(APORTES.cabecera[0])] = "aportes"
    por_cabecera[normalizar(FUENTES.cabecera[0])] = "fuentes"
    por_cabecera[normalizar(DATOS[0].etiqueta)] = "datos"
    tablas = {}
    for t in doc.tables:
        if not t.rows:
            continue
        clave = por_cabecera.get(normalizar(_primera_linea(t.rows[0].cells[0])))
        if clave and clave not in tablas:
            tablas[clave] = t
    esperadas = ["datos", "ubicacion", "comparabilidad", "instrumentos", "aportes", "subcategorias", "fuentes",
                 *(c.clave for c in CAJAS[1:])]
    faltan = [c for c in esperadas if c not in tablas]
    if faltan:
        raise PlantillaInvalida(
            "no parece una bitácora del laboratorio (v%d): faltan las tablas %s. ¿Se cambiaron los títulos?"
            % (VERSION, ", ".join(faltan))
        )
    return tablas


def _columnas_vigencia(t) -> dict[str, int]:
    """Columna de cada gobierno por su encabezado: «2018» marca el primero y «2026» el segundo, en
    cualquier parte del texto («Gobierno Duque (2018–2022)» sirve)."""
    cols = {}
    for i, celda in enumerate(t.rows[0].cells[1:], start=1):
        n = normalizar(_primera_linea(celda), numeracion=False)
        v = "2018-2022" if "2018" in n else "2022-2026" if "2026" in n else None
        if v and v not in cols:
            cols[v] = i
    if set(cols) != set(VIGENCIAS):
        raise PlantillaInvalida(
            f"en la tabla «{_primera_linea(t.rows[0].cells[0])}» no se reconocen las columnas de los dos "
            "gobiernos: sus encabezados deben mencionar 2018–2022 y 2022–2026"
        )
    return cols


def _columna(t, texto: str) -> int | None:
    for i, celda in enumerate(t.rows[0].cells):
        if normalizar(_primera_linea(celda)) == normalizar(texto):
            return i
    return None


def _id_politica(nombre: str, sector: str) -> str | None:
    dataset = DATOS_WEB / f"{sector}.json"
    if not nombre or not dataset.exists():
        return None
    n = normalizar(nombre)
    politicas = json.loads(dataset.read_text(encoding="utf-8")).get("politicas", [])
    for p in politicas:
        if normalizar(p["nombre"]) == n or p["id"] == slug(nombre):
            return p["id"]
    return None


def _nombre_instrumento(celda) -> str:
    """Lo que el grupo escribió en la celda de la etiqueta, debajo de «Instrumento N»; si lo escribió
    a continuación de la pista («Nombre: Ondas»), se quita la pista."""
    nombre = []
    for l in _texto_celda(celda).split("\n")[1:]:
        l = l.strip()
        if normalizar(l, numeracion=False).startswith("nombre"):
            l = l.split(":", 1)[1].strip() if ":" in l else ""
        if l:
            nombre.append(l)
    return "\n".join(nombre)


def _separadas(fila, cols) -> None:
    """Rechaza filas con celdas combinadas entre columnas que deben ir separadas."""
    tcs = [fila.cells[i]._tc for i in cols if i < len(fila.cells)]
    if len(set(map(id, tcs))) < len(tcs):
        raise PlantillaInvalida(
            f"la fila «{_primera_linea(fila.cells[0])}» tiene celdas combinadas: cada gobierno va en su columna"
        )


def leer(doc: Document, sector: str) -> dict:
    tablas = _indexar(doc)
    sin_llenar: list[str] = []
    notas: dict[str, str] = {}

    def celda(ruta: str, texto: str, *, obligatoria: bool = True):
        """Registra la celda: (presente, valor). Vacía -> sin_llenar (si es obligatoria);
        «Sin dato: nota» -> notas."""
        presente, valor, nota = _valor(texto)
        if not presente and obligatoria:
            sin_llenar.append(ruta)
        if nota:
            notas[ruta] = nota
        return presente, valor

    datos: dict[str, str] = {}
    for fila in tablas["datos"].rows:
        clave = _fila_de(DATOS, _primera_linea(fila.cells[0]))
        if clave and clave not in datos:
            datos[clave] = _texto_celda(fila.cells[1]) if len(fila.cells) > 1 and fila.cells[1]._tc is not fila.cells[0]._tc else ""
    for d in DATOS:
        if d.clave not in datos:
            raise PlantillaInvalida(f"falta la fila «{d.etiqueta}» en la tabla de datos del grupo")
    valores = {}
    for d in DATOS:
        presente, valor = celda(d.clave, datos[d.clave])
        valores[d.clave] = valor if presente and valor is not None else ""
    integrantes = [x.strip() for x in re.split(r"[,;\n]", valores["integrantes"]) if x.strip()]
    politica = {"nombre": valores["politica"]}
    if (pid := _id_politica(politica["nombre"], sector)) is not None:
        politica["id"] = pid
    grupo = {"nombre": valores["grupo"], "integrantes": integrantes}
    grupo |= {k: valores[k] for k in ("motivo", "expectativa") if valores[k]}
    b: dict = {"version": VERSION, "sector": sector, "grupo": grupo, "politica": politica}

    for c in COMPARATIVAS:
        t = tablas[c.clave]
        cols = _columnas_vigencia(t)
        b[c.clave] = {v: {} for v in VIGENCIAS}
        vistas = set()
        for fila in t.rows[1:]:
            clave = _fila_de(c.filas, _primera_linea(fila.cells[0]))
            if clave is None or clave in vistas:
                continue  # fila agregada por el grupo, vacía o repetida
            vistas.add(clave)
            celdas = {v: fila.cells[i] if i < len(fila.cells) else None for v, i in cols.items()}
            tcs = [x._tc for x in celdas.values() if x is not None]
            if len(set(map(id, tcs))) < len(tcs) or any(x is fila.cells[0]._tc for x in tcs):
                etiqueta = next(f.etiqueta for f in c.filas if f.clave == clave)
                raise PlantillaInvalida(
                    f"en «{c.cabecera}», la fila «{etiqueta}» tiene celdas combinadas: cada gobierno va en su columna"
                )
            for v, x in celdas.items():
                presente, valor = celda(f"{c.clave}.{v}.{clave}", _texto_celda(x) if x is not None else "")
                if presente:
                    b[c.clave][v][clave] = valor
        for f in c.filas:
            if f.clave not in vistas:
                raise PlantillaInvalida(f"falta la fila «{f.etiqueta}» en la tabla «{c.cabecera}»")

    # Instrumentos: solo las filas con algo escrito (son «hasta tres»; vacías no cuentan como faltantes).
    t = tablas["instrumentos"]
    cols = _columnas_vigencia(t)
    col_rel = _columna(t, INSTRUMENTOS_RELEVANCIA)
    instrumentos = []
    for n, fila in enumerate(t.rows[1:], start=1):
        _separadas(fila, [0, *cols.values(), *([col_rel] if col_rel is not None else [])])
        ins: dict = {}
        textos = {"nombre": _nombre_instrumento(fila.cells[0])}
        textos |= {v: _texto_celda(fila.cells[i]) for v, i in cols.items() if i < len(fila.cells)}
        if col_rel is not None and col_rel < len(fila.cells):
            textos["relevancia"] = _texto_celda(fila.cells[col_rel])
        for k, texto in textos.items():
            presente, valor = celda(f"instrumentos.{n}.{k}", texto, obligatoria=False)
            if presente:
                ins[k] = valor
        if ins:
            instrumentos.append(ins)
    b["instrumentos"] = instrumentos

    # Aportes a la red: cada fila con algo escrito, con su tipo por la etiqueta.
    aportes = []
    for fila in tablas["aportes"].rows[1:]:
        textos = [_texto_celda(x) for x in fila.cells[1 : len(APORTES.claves)]]
        if not any(textos):
            continue
        etiqueta = _primera_linea(fila.cells[0])
        aporte = {"tipo": _fila_de(TIPOS_APORTE, etiqueta) or normalizar(etiqueta) or "otro"}
        aporte |= {k: t for k, t in zip(APORTES.claves[1:], textos) if t}
        aportes.append(aporte)
    b["aportes"] = aportes

    for c in CAJAS:
        filas = tablas[c.clave].rows
        presente, valor = celda(c.clave, "\n".join(_texto_celda(f.cells[0]) for f in filas[1:]).strip())
        if presente:
            b[c.clave] = valor

    fuentes = []
    for fila in tablas["fuentes"].rows[1:]:
        textos = [_texto_celda(x) for x in fila.cells[: len(FUENTES.claves)]]
        if any(textos):  # sin inventar valores: solo lo que el grupo escribió
            fuentes.append({k: t for k, t in zip(FUENTES.claves, textos) if t})
    b["fuentes"] = fuentes
    if notas:
        b["notas"] = notas
    b["sin_llenar"] = sin_llenar
    return b


def validar(b: dict) -> list[str]:
    schema = json.loads(SCHEMA.read_text(encoding="utf-8"))
    return [
        f"{'/'.join(str(x) for x in e.absolute_path) or '(raíz)'}: {e.message}"
        for e in sorted(Draft202012Validator(schema).iter_errors(b), key=str)
    ]


# ─────────────────────────── Probar (ida y vuelta) ───────────────────────────


def probar() -> list[str]:
    """Genera la plantilla, la llena con el ejemplo CTeI, la lee y compara. Devuelve errores."""
    errores: list[str] = []
    ejemplo = json.loads(EJEMPLO_JSON.read_text(encoding="utf-8"))
    errores += [f"ejemplo: {e}" for e in validar(ejemplo)]
    with tempfile.TemporaryDirectory() as tmp:
        ruta = Path(tmp) / "b.docx"
        llenar(plantilla(), ejemplo).save(ruta)
        leida = leer(Document(ruta), ejemplo["sector"])
        vacia_ruta = Path(tmp) / "v.docx"
        plantilla().save(vacia_ruta)
        vacia = leer(Document(vacia_ruta), ejemplo["sector"])
    errores += [f"leída: {e}" for e in validar(leida)]
    if leida != ejemplo:
        distintas = sorted(k for k in set(leida) | set(ejemplo) if leida.get(k) != ejemplo.get(k))
        errores.append(f"ida y vuelta: difiere en {', '.join(distintas)}")
    if any(vacia[c.clave][v] for c in COMPARATIVAS for v in VIGENCIAS) or any(
        vacia[k] for k in ("instrumentos", "aportes", "fuentes")
    ):
        errores.append("la plantilla vacía no debería tener contenido")
    esperadas = len(DATOS) + sum(len(c.filas) * len(VIGENCIAS) for c in COMPARATIVAS) + len(CAJAS)
    if len(vacia["sin_llenar"]) != esperadas:
        errores.append(f"plantilla vacía: {len(vacia['sin_llenar'])} campos sin llenar, se esperaban {esperadas}")
    errores += _probar_tolerancia(ejemplo)
    return errores


def _probar_tolerancia(ejemplo: dict) -> list[str]:
    """Ediciones típicas de Word/Google Docs que el lector debe tolerar, y plantillas que debe rechazar."""
    errores = []
    doc = llenar(plantilla(), ejemplo)
    sub = _indexar(doc)["subcategorias"]
    # 1) etiquetas en mayúsculas y sin tildes; 2) fila extra agregada por el grupo; 3) «sin dato.» en minúsculas
    fila_obj = next(f for f in sub.rows if normalizar(_primera_linea(f.cells[0])) == "objetivo")
    fila_obj.cells[0].paragraphs[0].text = "4.1 OBJETIVO"
    fila_pob = next(f for f in sub.rows if normalizar(_primera_linea(f.cells[0])) == "poblacion")
    fila_pob.cells[0].paragraphs[0].text = "4.3 Poblacion"
    fila_pob.cells[1].paragraphs[0].text = "sin dato."
    extra = sub.add_row()
    extra.cells[0].paragraphs[0].text = "Otra cosa que el grupo quiso anotar"
    with tempfile.TemporaryDirectory() as tmp:
        ruta = Path(tmp) / "t.docx"
        doc.save(ruta)
        leida = leer(Document(ruta), ejemplo["sector"])
    if leida["subcategorias"] != ejemplo["subcategorias"]:
        errores.append("tolerancia: etiquetas en mayúsculas / sin tildes o fila extra rompen la lectura")
    # Una plantilla ajena se rechaza con un mensaje claro.
    try:
        otro = Document()
        otro.add_table(rows=1, cols=1).rows[0].cells[0].paragraphs[0].text = "Tabla cualquiera"
        leer(otro, ejemplo["sector"])
        errores.append("tolerancia: un documento ajeno debería rechazarse")
    except PlantillaInvalida:
        pass
    errores += _probar_casos_reales(ejemplo)
    return errores


def _fila(tabla, etiqueta: str):
    return next(f for f in tabla.rows if normalizar(_primera_linea(f.cells[0])) == normalizar(etiqueta))


def _leer_editada(ejemplo: dict, editar) -> dict:
    doc = llenar(plantilla(), ejemplo)
    editar(_indexar(doc))
    with tempfile.TemporaryDirectory() as tmp:
        ruta = Path(tmp) / "e.docx"
        doc.save(ruta)
        return leer(Document(ruta), ejemplo["sector"])


def _rechaza(ejemplo: dict, editar) -> bool:
    try:
        _leer_editada(ejemplo, editar)
        return False
    except PlantillaInvalida:
        return True


def _probar_casos_reales(ejemplo: dict) -> list[str]:
    """Casos de edición real que antes perdían o torcían datos en silencio."""
    errores = []

    # Columnas de gobierno combinadas: se rechaza (antes duplicaba el dato en las dos vigencias).
    def combinar(t):
        f = _fila(t["subcategorias"], "4.6 Metas")
        f.cells[1].merge(f.cells[2])
    if not _rechaza(ejemplo, combinar):
        errores.append("casos: celdas de gobierno combinadas deberían rechazarse")

    # Fila de datos borrada o renombrada: se rechaza (antes perdía lo escrito).
    def borrar_integrantes(t):
        f = _fila(t["datos"], "Integrantes")
        f._tr.getparent().remove(f._tr)
    if not _rechaza(ejemplo, borrar_integrantes):
        errores.append("casos: falta de la fila «Integrantes» debería rechazarse")

    # Variantes de «Sin dato», con nota: null + la nota en `notas`.
    def variantes(t):
        _fila(t["subcategorias"], "4.5 Recursos").cells[1].paragraphs[0].text = "S/D"
        _fila(t["subcategorias"], "4.2 Instituciones").cells[1].paragraphs[0].text = "Sin dato: el balance no nombra entidades"
    leida = _leer_editada(ejemplo, variantes)
    s18 = leida["subcategorias"]["2018-2022"]
    if s18.get("recursos", "x") is not None or s18.get("instituciones", "x") is not None:
        errores.append("casos: «S/D» y «Sin dato: nota» deberían leerse como hueco (null)")
    if leida.get("notas", {}).get("subcategorias.2018-2022.instituciones") != "el balance no nombra entidades":
        errores.append("casos: la nota de un «Sin dato: …» debería guardarse en `notas`")

    # Tabla anidada dentro de una celda: su texto se conserva.
    def anidar(t):
        celda = _fila(t["subcategorias"], "4.1 Objetivo").cells[2]
        celda.add_table(rows=1, cols=1).rows[0].cells[0].paragraphs[0].text = "texto anidado"
    if "texto anidado" not in (_leer_editada(ejemplo, anidar)["subcategorias"]["2022-2026"].get("objetivo") or ""):
        errores.append("casos: el texto de una tabla anidada se pierde")

    # Encabezado de gobierno con el año en medio, y etiqueta con texto agregado en la misma línea.
    def reescribir(t):
        t["ubicacion"].rows[0].cells[1].paragraphs[0].text = "Gobierno Duque (2018–2022)"
        _fila(t["subcategorias"], "4.1 Objetivo").cells[0].paragraphs[0].text = "4.1 Objetivo del gobierno"
    try:
        leida = _leer_editada(ejemplo, reescribir)
        if leida["ubicacion"] != ejemplo["ubicacion"] or leida["subcategorias"] != ejemplo["subcategorias"]:
            errores.append("casos: encabezado «Gobierno Duque (2018–2022)» o etiqueta ampliada leen mal")
    except PlantillaInvalida as e:
        errores.append(f"casos: encabezado con el año en medio o etiqueta ampliada se rechazan ({e})")

    # Datos vacíos: todos quedan registrados en `sin_llenar`.
    def vaciar(t):
        for f in t["datos"].rows:
            for p in f.cells[1].paragraphs:
                p.text = ""
    faltan = set(_leer_editada(ejemplo, vaciar)["sin_llenar"])
    if not {d.clave for d in DATOS} <= faltan:
        errores.append(f"casos: los datos vacíos del grupo deberían ir a sin_llenar (hay {sorted(faltan)})")

    # Nombre del instrumento escrito a continuación de la pista, en el mismo párrafo.
    def nombre_en_pista(t):
        celda = t["instrumentos"].rows[1].cells[0]
        for p in celda.paragraphs[1:]:
            p._p.getparent().remove(p._p)
        celda.add_paragraph("Nombre: Programa Ondas")
    if _leer_editada(ejemplo, nombre_en_pista)["instrumentos"][0].get("nombre") != "Programa Ondas":
        errores.append("casos: «Nombre: X» en la celda del instrumento debería leerse como su nombre")

    # Dos aportes del mismo tipo (fila agregada por el grupo): se leen los dos.
    def otro_aporte(t):
        f = t["aportes"].add_row()
        f.cells[0].paragraphs[0].text = "Añadir"
        f.cells[2].paragraphs[0].text = "Otro instrumento que falta"
    leida = _leer_editada(ejemplo, otro_aporte)
    if len(leida["aportes"]) != len(ejemplo["aportes"]) + 1 or leida["aportes"][-1]["tipo"] != "anadir":
        errores.append("casos: un aporte en una fila agregada debería leerse con su tipo")

    # Columnas de gobierno combinadas en un instrumento: se rechaza.
    def combinar_instrumento(t):
        f = t["instrumentos"].rows[2]
        f.cells[1].merge(f.cells[2])
    if not _rechaza(ejemplo, combinar_instrumento):
        errores.append("casos: celdas de gobierno combinadas en un instrumento deberían rechazarse")
    return errores


def publicadas_al_dia() -> list[str]:
    """Las .docx publicadas en web/public se leen con la estructura actual."""
    errores = []
    if not PLANTILLA.exists():
        return [f"falta {PLANTILLA}: uv run extraccion/bitacora.py generar"]
    try:
        vacia = leer(Document(PLANTILLA), "ciencia-tecnologia")
        if any(vacia[k] for k in ("instrumentos", "aportes", "fuentes")) or any(
            vacia[c.clave][v] for c in COMPARATIVAS for v in VIGENCIAS
        ):
            errores.append(f"{PLANTILLA} no está vacía")
    except PlantillaInvalida as e:
        errores.append(f"{PLANTILLA} desactualizada ({e}): regenerarla")
    if EJEMPLO_DOCX.exists():
        ejemplo = json.loads(EJEMPLO_JSON.read_text(encoding="utf-8"))
        try:
            if leer(Document(EJEMPLO_DOCX), ejemplo["sector"]) != ejemplo:
                errores.append(f"{EJEMPLO_DOCX} no coincide con {EJEMPLO_JSON.name}: regenerarla")
        except PlantillaInvalida as e:
            errores.append(f"{EJEMPLO_DOCX} desactualizada ({e})")
    return errores


# ─────────────────────────── CLI ───────────────────────────


def main() -> int:
    ap = argparse.ArgumentParser(description=__doc__.split("\n\n")[0], formatter_class=argparse.RawDescriptionHelpFormatter)
    sub = ap.add_subparsers(dest="cmd", required=True)
    g = sub.add_parser("generar", help="plantilla vacía")
    g.add_argument("--salida", type=Path, default=PLANTILLA)
    e = sub.add_parser("ejemplo", help="plantilla llena con el ejemplo CTeI")
    e.add_argument("--salida", type=Path, default=EJEMPLO_DOCX)
    l = sub.add_parser("leer", help=".docx llena -> JSON validado")
    l.add_argument("docx", type=Path)
    l.add_argument("--slug", required=True, help="sector, p. ej. ciencia-tecnologia")
    l.add_argument("--salida", type=Path, help="por defecto data/bitacoras/<slug>/<grupo>.json")
    l.add_argument("--forzar", action="store_true", help="sobrescribir si ya existe")
    sub.add_parser("probar", help="ida y vuelta con el ejemplo CTeI")
    a = ap.parse_args()

    if a.cmd in ("generar", "ejemplo"):
        doc = plantilla()
        if a.cmd == "ejemplo":
            llenar(doc, json.loads(EJEMPLO_JSON.read_text(encoding="utf-8")))
        a.salida.parent.mkdir(parents=True, exist_ok=True)
        doc.save(a.salida)
        print(f"ok  {a.salida}")
        return 0
    if a.cmd == "probar":
        errores = probar() + publicadas_al_dia()
        for err in errores:
            print(f"ERROR {err}", file=sys.stderr)
        if not errores:
            print("ok  bitácora: ida y vuelta con el ejemplo CTeI")
        return 1 if errores else 0

    try:
        b = leer(Document(a.docx), a.slug)
    except PlantillaInvalida as err:
        print(f"ERROR {a.docx}: {err}", file=sys.stderr)
        return 1
    errores = validar(b)
    for err in errores:
        print(f"ERROR {err}", file=sys.stderr)
    if errores:
        return 1
    salida = a.salida or SALIDA / a.slug / f"{slug(b['grupo']['nombre'] or b['politica']['nombre'] or a.docx.stem)}.json"
    if salida.exists() and not a.forzar:
        print(f"ERROR {salida} ya existe (¿otro grupo con el mismo nombre?): use --salida o --forzar", file=sys.stderr)
        return 1
    salida.parent.mkdir(parents=True, exist_ok=True)
    salida.write_text(json.dumps(b, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")
    huecos = sum(1 for c in COMPARATIVAS for v in VIGENCIAS for x in b[c.clave][v].values() if x is None)
    huecos += sum(1 for ins in b["instrumentos"] for x in ins.values() if x is None)
    print(f"ok  {salida} · política: {b['politica'].get('id', b['politica']['nombre'] + ' (sin área identificada)')}"
          f" · huecos declarados: {huecos} · sin llenar: {len(b['sin_llenar'])}")
    return 0


if __name__ == "__main__":
    sys.exit(main())
