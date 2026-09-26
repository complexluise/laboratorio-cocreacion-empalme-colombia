# /// script
# requires-python = ">=3.11"
# dependencies = ["pymupdf>=1.24", "openpyxl>=3.1", "xlrd>=2.0"]
# ///
"""
Conversor de documentos de empalme a markdown.

Recorre `descargas/` (informes PDF) y `extraido/` (anexos: pdf/xlsx/xls) y produce
un espejo en `markdown/`, preservando la ruta periodo/sector/entidad.

- PDF  -> markdown de texto por página (PyMuPDF/fitz, extracción rápida y robusta;
  NO usa el modelo de layout que cuelga sobre PDFs escaneados). Detecta PDFs sin
  texto (escaneados) y los marca para eventual OCR, en vez de atascarse.
- xlsx/xls -> resumen markdown: por hoja, dimensiones + encabezados + muestra de filas.
  (No vuelca la tabla completa: las bases de beneficiarios/presupuesto son enormes;
   el resumen basta para el ejercicio de ideas y deja rastro de qué contiene cada anexo.)

Uso:
  uv run empalme_to_markdown.py                 # convierte descargas/ + extraido/
  uv run empalme_to_markdown.py --solo-informes # solo los PDF principales de descargas/
"""
from __future__ import annotations

import argparse
import sys
import warnings
from pathlib import Path

warnings.filterwarnings("ignore")

# La consola de Windows es cp1252 y revienta con tildes combinadas en los nombres.
for _s in (sys.stdout, sys.stderr):
    try:
        _s.reconfigure(encoding="utf-8", errors="replace")
    except Exception:
        pass


def _safe(s: str) -> str:
    """Nombre imprimible en cualquier consola."""
    return s.encode("ascii", "replace").decode()

FUENTES = ["descargas", "extraido"]
DEST = Path("markdown")
MUESTRA_FILAS = 15  # filas de ejemplo por hoja de cálculo


def rel_destino(src: Path, raiz: str) -> Path:
    rel = src.relative_to(raiz)
    return DEST / rel.with_suffix(rel.suffix + ".md")


def convertir_pdf(src: Path, dst: Path) -> str:
    import fitz  # PyMuPDF
    doc = fitz.open(str(src))
    partes = []
    chars = 0
    for i, page in enumerate(doc, 1):
        txt = page.get_text("text").strip()
        chars += len(txt)
        if txt:
            partes.append(f"\n\n## Página {i}\n\n{txt}")
    n_pag = doc.page_count
    doc.close()
    escaneado = chars < 20 * n_pag  # casi sin texto -> probablemente escaneado (imagen)
    tipo = "PDF (escaneado — sin texto extraíble, requiere OCR)" if escaneado else "PDF"
    encabezado = f"# {src.stem}\n\n> Fuente: `{src}`\n> Tipo: {tipo} · {n_pag} páginas\n\n---\n"
    dst.parent.mkdir(parents=True, exist_ok=True)
    dst.write_text(encabezado + "".join(partes), encoding="utf-8")
    return f"ESCANEADO ({n_pag}p, {chars} chars)" if escaneado else f"OK ({n_pag}p, {chars} chars)"


def _celda(v) -> str:
    if v is None:
        return ""
    return str(v).replace("\n", " ").replace("|", "\\|").strip()


def convertir_xlsx(src: Path, dst: Path) -> str:
    import openpyxl
    wb = openpyxl.load_workbook(src, read_only=True, data_only=True)
    partes = [f"# {src.stem}\n", f"> Fuente: `{src}`", f"> Tipo: hoja de cálculo · {len(wb.sheetnames)} hoja(s)\n", "---\n"]
    for hoja in wb.sheetnames:
        ws = wb[hoja]
        filas = ws.iter_rows(values_only=True)
        try:
            encab = next(filas)
        except StopIteration:
            partes.append(f"## Hoja: {hoja}\n\n_(vacía)_\n")
            continue
        cols = [_celda(c) for c in encab]
        muestra = []
        n = 0
        for r in filas:
            n += 1
            if len(muestra) < MUESTRA_FILAS:
                muestra.append([_celda(c) for c in r])
        partes.append(f"## Hoja: {hoja}\n")
        partes.append(f"- Columnas: {len([c for c in cols if c])}")
        partes.append(f"- Filas de datos (aprox.): {n}\n")
        if any(cols):
            ancho = len(cols)
            partes.append("| " + " | ".join(cols) + " |")
            partes.append("| " + " | ".join(["---"] * ancho) + " |")
            for r in muestra:
                r = (r + [""] * ancho)[:ancho]
                partes.append("| " + " | ".join(r) + " |")
            if n > MUESTRA_FILAS:
                partes.append(f"\n_… {n - MUESTRA_FILAS} filas más no mostradas._")
        partes.append("")
    wb.close()
    dst.parent.mkdir(parents=True, exist_ok=True)
    dst.write_text("\n".join(partes), encoding="utf-8")
    return f"OK ({len(wb.sheetnames)} hojas)"


def convertir_xls(src: Path, dst: Path) -> str:
    import xlrd
    wb = xlrd.open_workbook(src)
    partes = [f"# {src.stem}\n", f"> Fuente: `{src}`", f"> Tipo: hoja de cálculo (xls) · {wb.nsheets} hoja(s)\n", "---\n"]
    for sh in wb.sheets():
        partes.append(f"## Hoja: {sh.name}\n")
        partes.append(f"- Dimensiones: {sh.nrows} filas × {sh.ncols} columnas\n")
        if sh.nrows == 0:
            partes.append("_(vacía)_\n")
            continue
        cols = [_celda(sh.cell_value(0, c)) for c in range(sh.ncols)]
        partes.append("| " + " | ".join(cols) + " |")
        partes.append("| " + " | ".join(["---"] * sh.ncols) + " |")
        for r in range(1, min(sh.nrows, 1 + MUESTRA_FILAS)):
            partes.append("| " + " | ".join(_celda(sh.cell_value(r, c)) for c in range(sh.ncols)) + " |")
        if sh.nrows > 1 + MUESTRA_FILAS:
            partes.append(f"\n_… {sh.nrows - 1 - MUESTRA_FILAS} filas más no mostradas._")
        partes.append("")
    dst.parent.mkdir(parents=True, exist_ok=True)
    dst.write_text("\n".join(partes), encoding="utf-8")
    return f"OK ({wb.nsheets} hojas)"


CONVERTIDORES = {".pdf": convertir_pdf, ".xlsx": convertir_xlsx, ".xls": convertir_xls}


def main():
    p = argparse.ArgumentParser()
    p.add_argument("--solo-informes", action="store_true", help="solo PDFs en descargas/")
    args = p.parse_args()
    fuentes = ["descargas"] if args.solo_informes else FUENTES

    total = ok = err = skip = 0
    for raiz in fuentes:
        base = Path(raiz)
        if not base.exists():
            continue
        for src in sorted(base.rglob("*")):
            if not src.is_file():
                continue
            conv = CONVERTIDORES.get(src.suffix.lower())
            if not conv:
                continue
            total += 1
            dst = rel_destino(src, raiz)
            if dst.exists() and dst.stat().st_size > 0:
                skip += 1
                continue
            try:
                res = conv(src, dst)
                ok += 1
                print(f"  [{ok:>3}] {_safe(src.name)[:60]:60} -> {res}", flush=True)
            except Exception as e:  # noqa: BLE401
                err += 1
                print(f"  ERR  {_safe(src.name)[:60]:60} -> {_safe(str(e))}", file=sys.stderr, flush=True)
    print(f"\nTotal: {total} | convertidos: {ok} | saltados: {skip} | errores: {err}")


if __name__ == "__main__":
    main()
